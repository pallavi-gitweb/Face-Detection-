import csv
import time
from datetime import datetime
from pathlib import Path

import cv2
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter


# Reports are saved beside this Python file.
BASE_DIR = Path(__file__).resolve().parent
IMAGE_REPORT_DIR = BASE_DIR / "Image Reports"
LIVE_REPORT_DIR = BASE_DIR / "Live Camera Reports"


def timestamp_now():
    return datetime.now().astimezone().isoformat(timespec="seconds")


def safe_filename(text):
    """Return a filename-friendly version of text."""
    return "".join(c if c.isalnum() or c in "-_." else "_" for c in text)


def load_face_detector():
    cascade_path = cv2.data.haarcascades + "haarcascade_frontalface_default.xml"
    detector = cv2.CascadeClassifier(cascade_path)
    if detector.empty():
        raise RuntimeError(
            "Could not load OpenCV Haar Cascade. Reinstall opencv-python and try again."
        )
    return detector


def detect_and_draw_faces(frame, detector):
    """Detect faces and draw bounding boxes. Returns annotated frame and face count."""
    gray = cv2.cvtColor(frame, cv2.COLOR_BGR2GRAY)
    faces = detector.detectMultiScale(
        gray, scaleFactor=1.1, minNeighbors=5, minSize=(30, 30)
    )

    for i, (x, y, w, h) in enumerate(faces, start=1):
        cv2.rectangle(frame, (x, y), (x + w, y + h), (0, 255, 0), 2)
        cv2.putText(
            frame, f"Face {i}", (x, max(y - 10, 20)),
            cv2.FONT_HERSHEY_SIMPLEX, 0.7, (0, 0, 255), 2
        )
    return frame, len(faces)


def style_worksheet(ws):
    """Apply readable headers and sensible column widths."""
    if ws.max_row >= 1:
        for cell in ws[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="24527A")
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions

    for column_cells in ws.columns:
        values = [len(str(c.value)) if c.value is not None else 0 for c in column_cells]
        width = min(max(max(values, default=10) + 2, 12), 42)
        ws.column_dimensions[get_column_letter(column_cells[0].column)].width = width


def write_excel(path, sheets):
    """Create an Excel workbook. sheets is [(sheet_name, headers, rows), ...]."""
    wb = Workbook()
    first = True
    for sheet_name, headers, rows in sheets:
        ws = wb.active if first else wb.create_sheet()
        first = False
        ws.title = sheet_name[:31]
        ws.append(headers)
        for row in rows:
            ws.append(list(row))
        style_worksheet(ws)
    wb.save(path)


def write_csv(path, headers, rows):
    with path.open("w", newline="", encoding="utf-8-sig") as f:
        writer = csv.writer(f)
        writer.writerow(headers)
        writer.writerows(rows)


def detect_faces_in_image(detector):
    IMAGE_REPORT_DIR.mkdir(parents=True, exist_ok=True)
    image_path = input(
        "\nEnter the full path to an image (or drag the image file into this terminal):\n> "
    ).strip().strip('"')

    if not image_path:
        print("No image path provided.")
        return

    path = Path(image_path).expanduser()
    if not path.is_file():
        print(f"File not found: {path}")
        return

    frame = cv2.imread(str(path))
    if frame is None:
        print("OpenCV could not read this file. Try a JPG or PNG image.")
        return

    detection_time = timestamp_now()
    result, count = detect_and_draw_faces(frame.copy(), detector)
    stamp = datetime.now().strftime("%Y%m%d_%H%M%S_%f")
    stem = safe_filename(path.stem)
    report_stem = f"{stem}_{stamp}"

    annotated_path = IMAGE_REPORT_DIR / f"{report_stem}_annotated.jpg"
    csv_path = IMAGE_REPORT_DIR / f"{report_stem}_report.csv"
    excel_path = IMAGE_REPORT_DIR / f"{report_stem}_report.xlsx"

    if not cv2.imwrite(str(annotated_path), result):
        print("Could not save the annotated image.")
        return

    headers = ["Image filename", "Detection timestamp", "Total faces detected", "Annotated image"]
    row = [path.name, detection_time, count, annotated_path.name]
    write_csv(csv_path, headers, [row])
    write_excel(excel_path, [("Image Summary", headers, [row])])

    print(f"\nImage: {path.name}")
    print(f"Detection timestamp: {detection_time}")
    print(f"Total faces detected: {count}")
    print(f"Annotated image: {annotated_path}")
    print(f"CSV report: {csv_path}")
    print(f"Excel report: {excel_path}")

    cv2.imshow("Face Detection - Image (press any key to close)", result)
    cv2.waitKey(0)
    cv2.destroyAllWindows()


def summarize_intervals(frame_rows, interval_seconds, session_duration):
    """Summarize frame observations in elapsed-time intervals."""
    if session_duration <= 0:
        interval_count = 1
    else:
        interval_count = max(1, int((session_duration + interval_seconds - 1e-9) // interval_seconds))
    summaries = []
    for i in range(interval_count):
        start = i * interval_seconds
        end = min((i + 1) * interval_seconds, session_duration)
        selected = [
            row for row in frame_rows
            if start <= row[0] < (i + 1) * interval_seconds
        ]
        frames_with_faces = sum(1 for row in selected if row[3] > 0)
        observations = sum(row[3] for row in selected)
        max_faces = max((row[3] for row in selected), default=0)
        summaries.append([
            f"{start:.1f}–{end:.1f} seconds",
            round(start, 3), round(end, 3),
            frames_with_faces, observations, max_faces
        ])
    return summaries


def detect_faces_from_webcam(detector):
    LIVE_REPORT_DIR.mkdir(parents=True, exist_ok=True)
    print("\nChoose interval for the live-camera report:")
    print("1. 10 seconds")
    print("2. 30 seconds")
    print("3. 60 seconds")
    interval_choice = input("Enter 1, 2, or 3 [default 1]: ").strip() or "1"
    interval_seconds = {"1": 10, "2": 30, "3": 60}.get(interval_choice)
    if interval_seconds is None:
        print("Invalid interval. Using 10 seconds.")
        interval_seconds = 10

    camera = cv2.VideoCapture(0)
    if not camera.isOpened():
        print(
            "Could not open the webcam. Check that it is connected, "
            "not being used by another app, and allowed in Windows privacy settings."
        )
        camera.release()
        return

    session_start = datetime.now().astimezone()
    start_clock = time.monotonic()
    frame_rows = []
    max_faces_in_frame = 0
    print("Webcam started. Press Q or Esc in the video window to stop and save reports.")

    try:
        while True:
            success, frame = camera.read()
            if not success:
                print("Could not read a frame from the webcam.")
                break

            elapsed = time.monotonic() - start_clock
            captured_at = timestamp_now()
            result, count = detect_and_draw_faces(frame, detector)
            max_faces_in_frame = max(max_faces_in_frame, count)

            # Each row represents one processed frame; repeated faces across frames are repeated observations.
            frame_rows.append([
                round(elapsed, 3), captured_at, len(frame_rows) + 1, count,
                count > 0, max_faces_in_frame
            ])

            cv2.putText(
                result, f"Faces detected: {count}", (10, 30),
                cv2.FONT_HERSHEY_SIMPLEX, 0.8, (255, 0, 0), 2
            )
            cv2.putText(
                result, f"Elapsed: {elapsed:.1f}s | Report interval: {interval_seconds}s",
                (10, 60), cv2.FONT_HERSHEY_SIMPLEX, 0.6, (255, 0, 0), 2
            )
            cv2.imshow("Live Face Detection - press Q to stop", result)
            key = cv2.waitKey(1) & 0xFF
            if key in (ord("q"), 27):
                break
    finally:
        session_duration = time.monotonic() - start_clock
        session_end = datetime.now().astimezone()
        camera.release()
        cv2.destroyAllWindows()

    stamp = session_start.strftime("%Y%m%d_%H%M%S")
    csv_path = LIVE_REPORT_DIR / f"live_camera_{stamp}_frame_log.csv"
    excel_path = LIVE_REPORT_DIR / f"live_camera_{stamp}_report.xlsx"

    frame_headers = [
        "Elapsed seconds", "Detection timestamp", "Frame number",
        "Faces detected in frame", "Frame contains faces", "Running maximum faces in a frame"
    ]
    frame_csv_rows = [
        [row[0], row[1], row[2], row[3], "Yes" if row[4] else "No", row[5]]
        for row in frame_rows
    ]
    write_csv(csv_path, frame_headers, frame_csv_rows)

    interval_rows = summarize_intervals(frame_rows, interval_seconds, session_duration)
    interval_headers = [
        "Time interval", "Start seconds", "End seconds", "Frames with faces",
        "Face observations", "Max faces in one frame"
    ]

    total_frames = len(frame_rows)
    frames_with_faces = sum(1 for row in frame_rows if row[3] > 0)
    total_observations = sum(row[3] for row in frame_rows)
    summary_headers = ["Metric", "Value"]
    summary_rows = [
        ["Session start time", session_start.isoformat(timespec="seconds")],
        ["Session end time", session_end.isoformat(timespec="seconds")],
        ["Session duration (seconds)", round(session_duration, 3)],
        ["Selected interval (seconds)", interval_seconds],
        ["Video frames processed", total_frames],
        ["Frames containing faces", frames_with_faces],
        ["Total face observations across frames", total_observations],
        ["Maximum faces visible in a single frame", max((r[3] for r in frame_rows), default=0)],
        ["Counting note", "Face observations, not unique people; repeated detections are expected."]
    ]
    write_excel(excel_path, [
        ("Session Summary", summary_headers, summary_rows),
        ("Interval Summary", interval_headers, interval_rows),
        ("Frame Log", frame_headers, frame_csv_rows),
    ])

    print("\n=== LIVE CAMERA REPORT SAVED ===")
    print(f"Start: {session_start.isoformat(timespec='seconds')}")
    print(f"End: {session_end.isoformat(timespec='seconds')}")
    print(f"Duration: {session_duration:.2f} seconds")
    print(f"Frames processed: {total_frames}")
    print(f"Frames containing faces: {frames_with_faces}")
    print(f"Total face observations: {total_observations}")
    print(f"Maximum faces in one frame: {max((r[3] for r in frame_rows), default=0)}")
    print(f"CSV frame log: {csv_path}")
    print(f"Excel workbook: {excel_path}")


def main():
    print("=" * 48)
    print("          OPENCV FACE DETECTION")
    print("=" * 48)
    print("1. Detect faces in an image")
    print("2. Detect faces using live webcam")
    choice = input("Choose 1 or 2: ").strip()

    detector = load_face_detector()
    if choice == "1":
        detect_faces_in_image(detector)
    elif choice == "2":
        detect_faces_from_webcam(detector)
    else:
        print("Invalid choice. Run the program again and choose 1 or 2.")


if __name__ == "__main__":
    main()
